"""One workbook contract, one calculation layer, no implicit sample fallbacks."""
from __future__ import annotations

import calendar
import hashlib
import math
from io import BytesIO
from datetime import date, datetime, timedelta, timezone
from pathlib import Path

from openpyxl import load_workbook


class WorkbookError(ValueError):
    pass


SCHEMA = {
    "Funds": "fund_id ticker name strategy description expense_ratio inception_date is_demo benchmark_label cash_weight".split(),
    "Daily Data": "fund_id date nav market_price benchmark_index net_assets shares_outstanding distribution_per_share".split(),
    "Holdings": "fund_id date ticker identifier name sector country currency security_type quantity price fx_rate market_value".split(),
    "Distributions": "fund_id ex_date record_date payable_date income short_term_gain long_term_gain return_of_capital".split(),
    "Documents": "fund_id document_type title url effective_date status".split(),
}


def fail(where, message):
    raise WorkbookError(f"{where}: {message}")


def number(value, where, *, minimum=None):
    if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value):
        fail(where, "expected a finite numeric cell, not text or a missing formula result")
    if minimum is not None and value < minimum:
        fail(where, f"must be at least {minimum}")
    return float(value)


def day(value, where):
    if isinstance(value, datetime):
        value = value.date()
    if isinstance(value, date):
        return value
    # ISO dates are accepted for interoperability; ambiguous localized dates are not.
    try:
        return date.fromisoformat(str(value))
    except ValueError:
        fail(where, "expected an Excel date or YYYY-MM-DD")


def text(value, where):
    if value is None or not str(value).strip():
        fail(where, "required value is empty")
    return str(value).strip()


def read_tables(path):
    raw = path if isinstance(path, bytes) else Path(path).read_bytes()
    values = load_workbook(BytesIO(raw), data_only=True, read_only=True)
    formulae = load_workbook(BytesIO(raw), data_only=False, read_only=True)
    tables = {}
    try:
        for sheet, fields in SCHEMA.items():
            if sheet not in values.sheetnames:
                fail(sheet, "required sheet is missing")
            rows = values[sheet].iter_rows(values_only=True)
            headers = list(next(rows, []))
            if headers != fields:
                fail(sheet, "header must be: " + ", ".join(fields))
            data = []
            for idx, row in enumerate(rows, 2):
                if not any(v is not None for v in row):
                    continue
                data.append(dict(zip(fields, row), _row=idx, _where=f"{sheet}!{idx}"))
            for row in formulae[sheet].iter_rows():
                for cell in row:
                    if cell.data_type == "f" and values[sheet][cell.coordinate].value is None:
                        fail(f"{sheet}!{cell.coordinate}", "formula has no saved result; recalculate and save in Excel")
            tables[sheet] = data
    finally:
        values.close()
        formulae.close()
    return tables


def months_before(d, months):
    index = d.year * 12 + d.month - 1 - months
    year, month = divmod(index, 12)
    month += 1
    return date(year, month, min(d.day, calendar.monthrange(year, month)[1]))


def performance(rows):
    """Illustrative close-to-close return series with ex-date distribution reinvestment.

    These are demo calculations, not a claim of standardized advertising compliance.
    Calendar anniversaries use the last available observation on/before the boundary.
    """
    out = []
    nav_index = market_index = 100.0
    for i, row in enumerate(rows):
        if i:
            previous = rows[i - 1]
            nav_index *= (row["nav"] + row["distribution_per_share"]) / previous["nav"]
            market_index *= (row["market_price"] + row["distribution_per_share"]) / previous["market_price"]
        out.append(dict(row, nav_total_return=nav_index, market_total_return=market_index,
                        premium_discount=row["market_price"] / row["nav"] - 1))
    last = out[-1]
    last_day = date.fromisoformat(last["date"])
    periods = {"1M": months_before(last_day, 1), "3M": months_before(last_day, 3),
               "6M": months_before(last_day, 6), "YTD": date(last_day.year - 1, 12, 31),
               "1Y": months_before(last_day, 12), "3Y": months_before(last_day, 36),
               "5Y": months_before(last_day, 60), "10Y": months_before(last_day, 120),
               "Since demo start": date.fromisoformat(out[0]["date"])}
    returns = {}
    for label, boundary in periods.items():
        available = [r for r in out if r["date"] <= boundary.isoformat()]
        base = available[-1] if available else None
        if base is None or base == last:
            returns[label] = None
            continue
        days = (last_day - date.fromisoformat(base["date"])).days
        annualized = label in ("3Y", "5Y", "10Y") or (label == "Since demo start" and days > 365)
        exponent = 365.25 / days if annualized else 1
        returns[label] = {key: (last[key] / base[key]) ** exponent - 1
                          for key in ("nav_total_return", "market_total_return", "benchmark_index")}
        returns[label].update(start_date=base["date"], annualized=annualized)
    return out, returns


def disclosure_periods(rows):
    """Last completed calendar year and completed subsequent quarters (as of dataset)."""
    as_of = date.fromisoformat(rows[-1]["date"])
    periods = [(str(as_of.year - 1), date(as_of.year - 1, 1, 1), date(as_of.year - 1, 12, 31))]
    for quarter in range(1, (as_of.month - 1) // 3 + 1):
        end_month = quarter * 3
        periods.append((f"{as_of.year} Q{quarter}", date(as_of.year, end_month - 2, 1),
                        date(as_of.year, end_month, calendar.monthrange(as_of.year, end_month)[1])))
    result = []
    for label, start, end in periods:
        window = [r for r in rows if start.isoformat() <= r["date"] <= end.isoformat()]
        if not window:
            continue
        result.append({"period": label, "start": window[0]["date"], "end": window[-1]["date"],
                       "observations": len(window),
                       "premium_days": sum(r["premium_discount"] > 1e-10 for r in window),
                       "discount_days": sum(r["premium_discount"] < -1e-10 for r in window),
                       "at_nav_days": sum(abs(r["premium_discount"]) <= 1e-10 for r in window)})
    return result


def import_workbook(path):
    path = Path(path)
    workbook_bytes = path.read_bytes()
    tables = read_tables(workbook_bytes)
    funds = {}
    for raw in tables["Funds"]:
        where = raw["_where"]
        fund_id = text(raw["fund_id"], where)
        if not fund_id.isascii() or not fund_id.isalnum() or fund_id.lower() != fund_id:
            fail(where, "fund_id must be lowercase ASCII letters/digits")
        if fund_id in funds:
            fail(where, f"duplicate fund_id {fund_id}")
        if raw["is_demo"] is not True:
            fail(where, "this review publisher only accepts is_demo=TRUE; live publishing needs a separate reviewed release")
        fund = {k: text(raw[k], f"{where} {k}") for k in ("ticker", "name", "strategy", "description", "benchmark_label")}
        fund.update(fund_id=fund_id, is_demo=True, inception_date=day(raw["inception_date"], where).isoformat(),
                    expense_ratio=number(raw["expense_ratio"], where, minimum=0),
                    cash_weight=number(raw["cash_weight"], where, minimum=0), daily=[], holdings=[], distributions=[], documents=[])
        if fund["expense_ratio"] >= 1 or fund["cash_weight"] > 1:
            fail(where, "rates must be decimal fractions: 0.0045 means 0.45%")
        funds[fund_id] = fund
    if not funds:
        fail("Funds", "at least one fund is required")
    seen = set()
    for raw in tables["Daily Data"]:
        where = raw["_where"]
        fund = funds.get(str(raw["fund_id"]))
        if fund is None:
            fail(where, "unknown fund_id")
        d = day(raw["date"], where)
        key = (fund["fund_id"], d)
        if key in seen:
            fail(where, "duplicate fund/date")
        seen.add(key)
        row = {k: number(raw[k], f"{where} {k}", minimum=0) for k in SCHEMA["Daily Data"][2:]}
        if any(row[k] <= 0 for k in ("nav", "market_price", "benchmark_index", "net_assets", "shares_outstanding")):
            fail(where, "prices, index, assets and shares must be positive")
        if abs(row["nav"] * row["shares_outstanding"] - row["net_assets"]) > max(0.02, row["net_assets"] * 1e-8):
            fail(where, "NAV × shares outstanding does not reconcile to net assets")
        if d.isoformat() < fund["inception_date"]:
            fail(where, "observation precedes demo inception_date")
        fund["daily"].append(dict(row, date=d.isoformat()))
    holding_keys = set()
    for raw in tables["Holdings"]:
        where = raw["_where"]
        fund = funds.get(str(raw["fund_id"]))
        if fund is None:
            fail(where, "unknown fund_id")
        row = {k: text(raw[k], f"{where} {k}") for k in SCHEMA["Holdings"][2:9]}
        row["date"] = day(raw["date"], where).isoformat()
        if row["security_type"] not in ("equity", "cash"):
            fail(where, "only equity and cash are supported")
        key = (fund["fund_id"], row["identifier"])
        if key in holding_keys:
            fail(where, "duplicate security identifier; use one consolidated position per security")
        holding_keys.add(key)
        for k in ("quantity", "price", "fx_rate", "market_value"):
            row[k] = number(raw[k], f"{where} {k}", minimum=0)
        if row["quantity"] <= 0 or row["price"] <= 0 or row["fx_rate"] <= 0:
            fail(where, "quantity, price and FX must be positive")
        if row["currency"] == "USD" and row["fx_rate"] != 1:
            fail(where, "USD FX must be 1")
        if abs(row["quantity"] * row["price"] * row["fx_rate"] - row["market_value"]) > 0.05:
            fail(where, "quantity × price × FX does not reconcile to market_value (USD)")
        fund["holdings"].append(row)
    for sheet, target in (("Distributions", "distributions"), ("Documents", "documents")):
        keys = set()
        for raw in tables[sheet]:
            where = raw["_where"]
            fund = funds.get(str(raw["fund_id"]))
            if fund is None:
                fail(where, "unknown fund_id")
            if sheet == "Distributions":
                row = {k: day(raw[k], f"{where} {k}").isoformat() for k in ("ex_date", "record_date", "payable_date")}
                if row["record_date"] < row["ex_date"] or row["payable_date"] < row["record_date"]:
                    fail(where, "distribution dates are out of order")
                for k in SCHEMA[sheet][4:]:
                    row[k] = number(raw[k], f"{where} {k}", minimum=0)
                row["total"] = sum(row[k] for k in SCHEMA[sheet][4:])
                key = (fund["fund_id"], row["ex_date"])
            else:
                row = {k: str(raw[k] or "").strip() for k in SCHEMA[sheet][1:]}
                if row["status"] not in ("unavailable", "draft"):
                    fail(where, "demo documents must be unavailable or draft")
                if not row["title"] or not row["document_type"]:
                    fail(where, "document title/type required")
                if row["url"] and not row["url"].startswith("https://"):
                    fail(where, "document URL must use https")
                if row["effective_date"]:
                    row["effective_date"] = day(raw["effective_date"], where).isoformat()
                key = (fund["fund_id"], row["document_type"])
            if key in keys:
                fail(where, "duplicate record")
            keys.add(key)
            fund[target].append(row)
    for fund in funds.values():
        fund["daily"].sort(key=lambda r: r["date"])
        if len(fund["daily"]) < 2 or not fund["holdings"]:
            fail(fund["fund_id"], "at least two daily observations and a full holdings snapshot are required")
        latest = fund["daily"][-1]
        if any(h["date"] != latest["date"] for h in fund["holdings"]):
            fail(fund["fund_id"], "holdings and latest daily data must have the same as-of date")
        assets = sum(h["market_value"] for h in fund["holdings"])
        if abs(assets - latest["net_assets"]) > max(0.10, latest["net_assets"] * 1e-8):
            fail(fund["fund_id"], "demo holdings including cash must reconcile to net assets; this simple demo has no other assets/liabilities")
        cash_weight = sum(h["market_value"] for h in fund["holdings"] if h["security_type"] == "cash") / assets
        if abs(cash_weight - fund["cash_weight"]) > 1e-6:
            fail(fund["fund_id"], "cash_weight disagrees with the holdings cash position")
        for h in fund["holdings"]:
            h["weight"] = h["market_value"] / assets
        fund["holdings"].sort(key=lambda h: -h["market_value"])
        daily_by_date = {r["date"]: r for r in fund["daily"]}
        distributions = {r["ex_date"]: r["total"] for r in fund["distributions"]}
        for d, row in daily_by_date.items():
            if abs(row["distribution_per_share"] - distributions.get(d, 0)) > 1e-8:
                fail(fund["fund_id"], f"distribution amount on {d} disagrees between Daily Data and Distributions")
        if any(d not in daily_by_date for d in distributions):
            fail(fund["fund_id"], "each distribution ex-date needs a daily data observation")
        fund["daily"], fund["returns"] = performance(fund["daily"])
        fund["as_of"] = latest["date"]
        fund["disclosure_periods"] = disclosure_periods(fund["daily"])
        fund["spread"] = {"median_30_day": None, "status": "Unavailable — requires reviewed intraday NBBO data"}
        fund["distributions"].sort(key=lambda r: r["ex_date"], reverse=True)
    return {"schema_version": 1, "is_demo": True,
            "source_sha256": hashlib.sha256(workbook_bytes).hexdigest(),
            "generated_at": datetime.now(timezone.utc).isoformat(),
            "source_workbook": path.name, "funds": list(funds.values())}
